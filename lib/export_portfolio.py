import pandas as pd
from pathlib import Path
import numpy as np
from openpyxl import load_workbook

def export_to_excel(final_weights, capital):
    input_path = (
                Path(__file__).resolve().parent.parent
                / "data"
                / "Hand_in_sheet.xlsx"
            )
    output_path = (
                Path(__file__).resolve().parent.parent
                / "outputs"
                / "Hand_in_sheet.xlsx"
            )
    df = pd.read_excel(
        input_path,
        header=19,
        usecols="A:F"
    )
    df = df.loc[df['PERMNO'].notna()].copy()
    df['weight'] = df["PERMNO"].map(final_weights).fillna(0)
    df['shares'] = np.floor(
        capital * df["weight"] 
        / df["Close price 31-12-2025"]
        ).astype(int)
    df['investment'] = df['shares'] * df['Close price 31-12-2025']
    df['actual_weight'] = df['investment'] / capital
    rf_investment = capital - df['investment'].sum()

    wb = load_workbook(input_path)
    ws = wb.worksheets[0]
    for i, shares in df['shares'].items():
        ws.cell(row=int(i) + 21, column=8, value=int(shares))
    ws["C12"] = capital
    ws["H21"] = float(rf_investment)
    ws["I16"] = "=SUM(I22:I526)"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    